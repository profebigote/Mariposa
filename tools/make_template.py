#!/usr/bin/env python3
"""Génère le classeur modèle tools/modele_mariposa.xlsx (6 onglets importables + un onglet LEEME).
Usage : python3 tools/make_template.py   (nécessite openpyxl)"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
HEAD = PatternFill('solid', fgColor='C8402B'); HF = Font(bold=True, color='FFFFFF')
NOTE = PatternFill('solid', fgColor='FBEFD9')

def opts(n):
    c = []
    for k in range(1, n + 1):
        c += [f'o{k}_es', f'o{k}_fr', f'o{k}_ok', f'o{k}_next', f'o{k}_set', f'o{k}_hint_es', f'o{k}_hint_fr', f'o{k}_why_es', f'o{k}_why_fr', f'o{k}_explain_es', f'o{k}_explain_fr', f'o{k}_kind']
    return c
def acc(n):
    c = []
    for k in range(1, n + 1): c += [f'a{k}_text', f'a{k}_es', f'a{k}_fr', f'a{k}_next', f'a{k}_set']
    return c

SHEETS = {
 'Messages_UI': {
  'cols': ['key', 'kind', 'es', 'fr', 'activity', 'errorKind', 'level'],
  'rows': [
   ['err1.demo.a', 'err1', '¡Casi! Vuelve a mirar la frase.', 'Presque ! Regarde encore la phrase.', 'fill;transform', 'grammar', ''],
   ['ok1.demo.a', 'ok1', '¡Exacto!', 'Exactement !', '', '', ''],
   ['mission.demo', 'mission', '¡Misión cumplida! Ya sabes {skill}.', 'Mission accomplie ! Tu sais {skill}.', '', '', ''],
  ]},
 'Fiches_modules': {
  'cols': ['id','type','title','rubrique','theme','level','duration','status','version','objective_main','objectives_secondary','outcome','lexicon','grammar','prerequisites','skills','intro_es','intro_fr','scene','cast_json','doc_json','completion_es','completion_fr','review','next','data_json'],
  'rows': [
   ['m-demo-aventura','adventure','Demo : una aventura','situaciones','Café','A1','4','brouillon','1','Pedir algo en una cafetería','saludar;pedir con educación','pedir una bebida','un café;por favor','Quiero + nombre','','pedir-cafeteria','Entras en un café.','Tu entres dans un café.','cafe','{"cam":{"name":"Camarero","skin":1,"hair":0,"shirt":0},"tu":{"name":"Tú","player":true,"skin":0,"hair":2,"shirt":1}}','','¡Misión cumplida! Ya sabes pedir algo en una cafetería.','Mission accomplie !','','m-demo-lexico',''],
   ['m-demo-leccion','lesson','Demo : una lección visual','gramatica','Gerundio','A2','5','brouillon','1','Decir lo que estoy haciendo','formar el gerundio','decir qué haces ahora','comer;hablar;vivir','estar + gerundio','','estar-gerundio','Vamos a ver cómo se forma el gerundio.','Voyons comment se forme le gérondif.','','','','','','','',''],
   ['m-demo-lexico','vocab','Demo : frutas','palabras','Frutas','A1','4','brouillon','1','Reconocer las frutas','','reconocer cinco frutas','la manzana;la pera','','','reconocer-alimentos','Cinco frutas para empezar.','Cinq fruits pour commencer.','','','','','','','',''],
  ]},
 'Scenario_demo': {
  'cols': ['module_id','node_id','type','speaker','es','fr','next','cond','else','prompt_es','prompt_fr','tiles','hint_es','hint_fr','explain_es','explain_fr','solution_es','solution_fr'] + opts(4) + acc(3) + ['data_json','notes'],
  'rows': 'scenario'},
 'Lecons_visuelles': {
  'cols': ['module_id','step','phase','kind','title_es','title_fr','prompt_es','prompt_fr','options','answers','tiles','items','pairs','points_es','points_fr','hint_es','hint_fr','explain_es','explain_fr','why_es','why_fr','img','data_json','notes'],
  'rows': 'lessons'},
 'Evolution_papillon': {
  'cols': ['stage','name_es','name_fr','threshold','desc_fr','message_es','message_fr'],
  'rows': [
   [1,'La pequeña oruga','La petite chenille',0,'Curieuse, elle découvre ses premiers mots.','Mariposa empieza su viaje contigo.','Mariposa commence son voyage avec toi.'],
   [2,'La oruga exploradora','La chenille exploratrice',6,'Elle grandit et gagne des couleurs.','¡Mariposa está creciendo y tiene colores nuevos!','Mariposa grandit et a de nouvelles couleurs !'],
   [3,'La crisálida','La chrysalide',16,'Elle se prépare : ce que tu as appris se consolide.','Mariposa descansa en su crisálida: lo aprendido se está consolidando.','Mariposa se repose dans sa chrysalide : ce que tu as appris se consolide.'],
   [4,'La mariposa joven','Le jeune papillon',30,'Ses premières ailes se déploient.','¡Mariposa ha salido de la crisálida!','Mariposa est sortie de sa chrysalide !'],
   [5,'La mariposa viajera','Le papillon voyageur',48,'Ses ailes se parent de motifs : elle part explorer le monde hispanophone.','¡Mariposa vuela lejos y tú contigo!','Mariposa s’envole loin, et toi avec elle !'],
  ]},
 'Lexique_demo': {
  'cols': ['module_id','id','es','fr','emoji','sentence_es','sentence_fr','answers','example_es','example_fr','tags'],
  'rows': [
   ['m-demo-lexico','manzana','la manzana','la pomme','🍎','Quiero comprar tres ________.','Je veux acheter trois ________.','manzanas','Me como una manzana.','Je mange une pomme.','fruta'],
   ['m-demo-lexico','pera','la pera','la poire','🍐','','','','',' ','fruta'],
   ['m-demo-lexico','uvas','las uvas','le raisin','🍇','','','','','','fruta'],
   ['m-demo-lexico','fresa','la fresa','la fraise','🍓','','','','','','fruta'],
   ['m-demo-lexico','platano','el plátano','la banane','🍌','Desayuno un ________ con leche.','Je prends une ________ avec du lait au petit-déjeuner.','plátano','','','fruta'],
  ]},
}

def row(cols, **kw):
    r = [''] * len(cols)
    for k, v in kw.items(): r[cols.index(k)] = v
    return r

def scenario_rows(cols):
    M = 'm-demo-aventura'; R = []
    R.append(row(cols, module_id=M, node_id='n1', type='say', speaker='cam', es='¡Hola! ¿Qué te pongo?', fr='Bonjour ! Qu’est-ce que je te sers ?', next='n2', notes='say = une réplique d’un personnage (speaker = clé de cast)'))
    R.append(row(cols, module_id=M, node_id='n2', type='choice', prompt_es='¿Qué respondes?', prompt_fr='Que réponds-tu ?',
        o1_es='Un café, por favor.', o1_fr='Un café, s’il vous plaît.', o1_ok='oui', o1_next='n3', o1_set='bebida=café', o1_why_es='Corta y educada.', o1_why_fr='Courte et polie.',
        o2_es='Quiero un té, por favor.', o2_fr='Je voudrais un thé, s’il vous plaît.', o2_ok='oui', o2_next='n3', o2_set='bebida=té',
        o3_es='Adiós.', o3_fr='Au revoir.', o3_ok='non', o3_kind='situation', o3_hint_es='Acabas de llegar.', o3_hint_fr='Tu viens d’arriver.',
        notes='choice = plusieurs options ; plusieurs peuvent être justes (ok=oui) et mener à des suites différentes'))
    R.append(row(cols, module_id=M, node_id='n3', type='branch', cond='bebida=café', next='n3c', **{'else': 'n3t'}, notes='branch = aiguillage automatique selon une variable (set)'))
    R.append(row(cols, module_id=M, node_id='n3c', type='say', speaker='cam', es='Marchando un café.', fr='Un café, c’est parti.', next='n4'))
    R.append(row(cols, module_id=M, node_id='n3t', type='say', speaker='cam', es='Marchando un té.', fr='Un thé, c’est parti.', next='n4'))
    R.append(row(cols, module_id=M, node_id='n4', type='build', prompt_es='Pregunta el precio.', prompt_fr='Demande le prix.', tiles='¿Cuánto;es;cuesta;por favor;?',
        a1_text='¿Cuánto es?', a1_next='n5', a2_text='¿Cuánto cuesta?', a2_next='n5', hint_es='Dos preguntas posibles.', hint_fr='Deux questions possibles.', notes='build = phrase à reconstruire ; chaque formulation acceptée = une colonne aN_text'))
    R.append(row(cols, module_id=M, node_id='n5', type='say', speaker='cam', es='Son dos euros.', fr='Ça fait deux euros.', next='fin'))
    R.append(row(cols, module_id=M, node_id='fin', type='end', es='¡Misión cumplida!', fr='Mission accomplie !', notes='end = fin de l’aventure (déclenche le bilan)'))
    return R

def lesson_rows(cols):
    M = 'm-demo-leccion'; R = []
    R.append(row(cols, module_id=M, step=1, kind='observe', title_es='Mira', title_fr='Observe', data_json='{"examples":[{"es":"Estoy comiendo.","fr":"Je suis en train de manger.","img":"🍽️","highlight":["comiendo"]}],"question":{"es":"¿Qué ves en el verbo?","fr":"Que vois-tu dans le verbe ?"}}', notes='observe, schema, formula, contrast, rule : listes complexes dans data_json'))
    R.append(row(cols, module_id=M, step=2, kind='choice', title_es='Elige', title_fr='Choisis', prompt_es='Marta ______ la cena.', prompt_fr='Marta prépare le dîner.', options='*está preparando;está preparar;está preparado', hint_es='Mira la terminación.', hint_fr='Regarde la terminaison.', explain_es='estar + -ando / -iendo.', explain_fr='estar + -ando / -iendo.', notes='options : séparées par ; — l’option juste commence par *'))
    R.append(row(cols, module_id=M, step=3, kind='fill', title_es='Completa', title_fr='Complète', prompt_es='Pablo está ______ (hablar) por teléfono.', prompt_fr='', answers='hablando', hint_es='-ar → -ando', hint_fr='-ar → -ando', notes='answers : réponses acceptées séparées par ;'))
    R.append(row(cols, module_id=M, step=4, kind='recap', points_es='estar + gerundio;-ar → -ando', points_fr='estar + gérondif;-ar → -ando'))
    return R

def build():
    wb = Workbook(); ws0 = wb.active; ws0.title = 'LEEME'
    txt = ['Modèle de classeur Proyecto Mariposa',
     '',
     'Chaque onglet (sauf LEEME) est importable depuis Atelier > Importer. Ne renomme pas les onglets ni les en-têtes de colonnes.',
     'Une ligne = un élément : un message, une fiche, un nœud d’aventure, un écran de leçon, une étape du papillon, un mot.',
     'Le lien entre onglets est la colonne module_id (= id dans Fiches_modules).',
     'Listes : sépare les valeurs par un point-virgule ( ; ).  Variables : cle=valeur;cle2=valeur2.  Oui/non : oui, non (ou 1/0).',
     'data_json : pour les structures complexes, colle un objet JSON ; il est fusionné dans la ligne et l’emporte sur les autres colonnes.',
     'Le texte « es » est en espagnol, « fr » en français (aide). Une réplique sans « fr » reste valable.',
     'Tout contenu importé est un BROUILLON : relis-le dans Atelier, puis passe-le à « validé » / « publié ».',
     'Détail complet : docs/FORMAT_MODULES.md']
    for i, t in enumerate(txt, 1): ws0.cell(i, 1, t)
    ws0['A1'].font = Font(bold=True, size=14); ws0.column_dimensions['A'].width = 140
    for name, d in SHEETS.items():
        ws = wb.create_sheet(name); cols = d['cols']
        ws.append(cols)
        rows = d['rows']
        if rows == 'scenario': rows = scenario_rows(cols)
        elif rows == 'lessons': rows = lesson_rows(cols)
        for r in rows: ws.append(r)
        for i, c in enumerate(cols, 1):
            cell = ws.cell(1, i); cell.fill = HEAD; cell.font = HF; cell.alignment = Alignment(vertical='center')
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(38, len(c) + 4))
        ws.freeze_panes = 'A2'
        if name in ('Scenario_demo', 'Lecons_visuelles'):
            ws.column_dimensions[get_column_letter(cols.index('es') + 1) if 'es' in cols else 'A'].width = 34
    out = os.path.join(HERE, 'modele_mariposa.xlsx'); wb.save(out); print('écrit', out)

if __name__ == '__main__': build()
